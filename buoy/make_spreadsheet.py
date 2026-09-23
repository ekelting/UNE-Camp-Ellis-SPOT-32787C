"""
make_spreadsheet.py - builds Camp_Ellis_Buoy_Summary.xlsx (opens in Excel or Google Sheets).
Called by update_buoy_report.py; you don't need to run it directly.
"""
import datetime as dt

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import buoy_core as bc

F = 'Arial'
NAVY, TEAL, SAND, FOAM, CORAL = '0B3C5D', '1D7A8C', 'F4EBD9', 'E6F4F1', 'D9534F'
H_FILL = PatternFill('solid', fgColor=NAVY)
SEC_FILL = PatternFill('solid', fgColor=TEAL)
ALT_FILL = PatternFill('solid', fgColor=FOAM)
IN_FILL = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='C9D6DF')
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def _font(**kw):
    kw.setdefault('name', F)
    kw.setdefault('size', 10)
    return Font(**kw)


def _title(ws, text, sub=None, width=10):
    ws['A1'] = text
    ws['A1'].font = _font(size=16, bold=True, color='FFFFFF')
    for c in range(1, width + 1):
        ws.cell(row=1, column=c).fill = H_FILL
        if sub:
            ws.cell(row=2, column=c).fill = PatternFill('solid', fgColor=SAND)
    ws.row_dimensions[1].height = 30
    if sub:
        ws['A2'] = sub
        ws['A2'].font = _font(italic=True, color='444444')


def _header(ws, row, headers, col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font = _font(bold=True, color='FFFFFF')
        c.fill = SEC_FILL
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = BOX
    ws.row_dimensions[row].height = 32


def _table(ws, row, df, fmts, col=1, band=True):
    """Write a dataframe body starting at row. fmts: list of number formats per column."""
    for r, rec in enumerate(df.itertuples(index=False), start=row):
        for j, v in enumerate(rec):
            if isinstance(v, float) and (np.isnan(v) or np.isinf(v)):
                v = None
            if isinstance(v, (np.floating,)):
                v = float(v)
            if isinstance(v, (np.integer,)):
                v = int(v)
            if isinstance(v, pd.Timestamp):
                v = v.tz_localize(None).to_pydatetime() if v.tzinfo else v.to_pydatetime()
            c = ws.cell(row=r, column=col + j, value=v)
            c.font = _font()
            c.border = BOX
            if fmts and j < len(fmts) and fmts[j]:
                c.number_format = fmts[j]
            if band and (r - row) % 2:
                c.fill = ALT_FILL
    return row + len(df)


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _section(ws, row, text, width=6):
    c = ws.cell(row=row, column=1, value=text)
    c.font = _font(size=12, bold=True, color=NAVY)
    for j in range(1, width + 1):
        ws.cell(row=row, column=j).border = Border(bottom=Side(style='medium', color=TEAL))
    return row + 1


def _line_chart(title, ytitle, w=26, h=9):
    ch = LineChart()
    ch.title, ch.y_axis.title = title, ytitle
    ch.width, ch.height = w, h
    ch.legend.position = 'b'
    return ch


def build_xlsx(S, path, SS=None):
    st, mo, ev, dy, hr, ro = S['stats'], S['monthly'], S['storms'], S['daily'], S['hourly'], S['rose']
    mo = mo[~mo['partial']].reset_index(drop=True)
    wb = Workbook()

    # ================================================================ Hourly ==
    wh = wb.active
    wh.title = 'Hourly'
    h = hr.copy()
    h.index = bc.local(h.index).tz_localize(None)
    h = h.reset_index().rename(columns={'t': 'time', 'index': 'time'})
    h = h[['time', 'hs', 'tp', 'te', 'dm', 'power_kw_m', 'wspd', 'pres']]
    h = h.dropna(subset=['hs', 'pres'], how='all')
    _title(wh, '⏱️ Hourly averages — the raw material for the Overview formulas',
           'Local time (America/New_York). Cleaned data (spikes removed). Blank = no data that hour.', 8)
    _header(wh, 3, ['Time (local)', 'Sig. wave height Hs (m)', 'Peak period Tp (s)', 'Energy period Te (s)',
                    'Mean direction (° from)', 'Wave power (kW/m)', 'Wind speed* (m/s)', 'Pressure (hPa)'])
    end_h = _table(wh, 4, h, ['yyyy-mm-dd hh:mm', '0.00', '0.0', '0.0', '0', '0.00', '0.0', '0.0'], band=False) - 1
    _widths(wh, [18, 12, 11, 11, 12, 12, 12, 11])
    wh.freeze_panes = 'B4'
    HR = lambda c: f"Hourly!${c}$4:${c}${end_h}"

    # ================================================================= Daily ==
    wd = wb.create_sheet('Daily')
    d = dy.reset_index()
    d['date'] = pd.to_datetime(d['date'])
    d = d[['date', 'hs_mean', 'hs_max', 'tp_mean', 'te_mean', 'power_mean', 'dir_deg', 'dir', 'wind_mean',
           'wind_max', 'pres_mean', 'pres_min', 'batt_mean', 'humid_mean', 'n_waves_obs']]
    _title(wd, '📈 Daily summary', 'One row per local day. Charts are to the right →', 15)
    _header(wd, 3, ['Date', 'Hs mean (m)', 'Hs max (m)', 'Tp mean (s)', 'Te mean (s)', 'Power mean (kW/m)',
                    'Direction (°)', 'Direction', 'Wind mean* (m/s)', 'Wind max* (m/s)', 'Pressure mean (hPa)',
                    'Pressure min (hPa)', 'Battery (V)', 'Hull humidity (%)', 'Wave readings'])
    end_d = _table(wd, 4, d, ['yyyy-mm-dd', '0.00', '0.00', '0.0', '0.0', '0.00', '0', '@', '0.0', '0.0',
                              '0.0', '0.0', '0.00', '0', '0'], band=False) - 1
    _widths(wd, [12, 9, 9, 9, 9, 10, 10, 9, 9, 9, 10, 10, 9, 9, 9])
    wd.freeze_panes = 'B4'

    ch = _line_chart('🌊 Daily wave height', 'Hs (m)')
    ch.add_data(Reference(wd, min_col=2, max_col=3, min_row=3, max_row=end_d), titles_from_data=True)
    ch.set_categories(Reference(wd, min_col=1, min_row=4, max_row=end_d))
    ch.x_axis.number_format = 'mmm yy'
    wd.add_chart(ch, 'Q3')
    ch2 = _line_chart('🌡️ Daily minimum barometric pressure', 'hPa')
    ch2.add_data(Reference(wd, min_col=12, max_col=12, min_row=3, max_row=end_d), titles_from_data=True)
    ch2.set_categories(Reference(wd, min_col=1, min_row=4, max_row=end_d))
    ch2.y_axis.scaling.min = 980
    ch2.x_axis.number_format = 'mmm yy'
    wd.add_chart(ch2, 'Q23')

    # =============================================================== Monthly ==
    wm = wb.create_sheet('Monthly')
    m = mo[['label', 'coverage_pct', 'n_obs', 'hs_mean', 'hs_median', 'hs_p90', 'hs_max', 'tp_mean', 'te_mean',
            'dir', 'power_mean', 'energy_mj_m', 'pct_over_1m', 'pct_over_2m', 'pct_calm', 'swell_share',
            'wind_mean', 'wind_max', 'pres_mean', 'pres_min', 'storms', 'storm_hours']].copy()
    for c in ['coverage_pct', 'pct_over_1m', 'pct_over_2m', 'pct_calm', 'swell_share']:
        m[c] = m[c] / 100
    _title(wm, '📅 Month by month', 'Months with < 50% data coverage are left out. Charts below the table ↓', 22)
    _header(wm, 3, ['Month', 'Data coverage', 'Wave readings', 'Hs mean (m)', 'Hs median (m)', 'Hs 90th pct (m)',
                    'Hs max (m)', 'Tp mean (s)', 'Te mean (s)', 'Main direction', 'Power mean (kW/m)',
                    'Energy (MJ/m)', 'Time Hs ≥ 1 m', 'Time Hs ≥ 2 m', 'Time calm (< 0.3 m)', 'Swell share of energy',
                    'Wind mean* (m/s)', 'Wind max* (m/s)', 'Pressure mean (hPa)', 'Pressure min (hPa)',
                    'Storm events', 'Storm hours'])
    end_m = _table(wm, 4, m, ['@', '0%', '#,##0', '0.00', '0.00', '0.00', '0.00', '0.0', '0.0', '@', '0.00', '#,##0',
                              '0.0%', '0.0%', '0%', '0%', '0.0', '0.0', '0.0', '0.0', '0', '0']) - 1
    tr = end_m + 1
    wm.cell(row=tr, column=1, value='All months').font = _font(bold=True)
    for col, fn in [(3, 'SUM'), (4, 'AVERAGE'), (7, 'MAX'), (12, 'SUM'), (20, 'MIN'), (21, 'SUM'), (22, 'SUM')]:
        L = get_column_letter(col)
        c = wm.cell(row=tr, column=col, value=f'={fn}({L}4:{L}{end_m})')
        c.font = _font(bold=True)
        c.number_format = wm.cell(row=4, column=col).number_format
        c.border = BOX
    wm.cell(row=tr + 1, column=1, value='(Hs mean in the total row is the average of monthly means.)').font = _font(italic=True, size=8, color='777777')
    _widths(wm, [11] + [10] * 21)
    wm.freeze_panes = 'B4'

    b = BarChart()
    b.type, b.grouping = 'col', 'clustered'
    b.title, b.y_axis.title = '📊 Wave height by month', 'Hs (m)'
    b.add_data(Reference(wm, min_col=4, max_col=4, min_row=3, max_row=end_m), titles_from_data=True)
    b.add_data(Reference(wm, min_col=6, max_col=7, min_row=3, max_row=end_m), titles_from_data=True)
    b.set_categories(Reference(wm, min_col=1, min_row=4, max_row=end_m))
    b.width, b.height = 20, 9
    b.legend.position = 'b'
    wm.add_chart(b, f'A{tr + 3}')
    b2 = BarChart()
    b2.type = 'col'
    b2.title, b2.y_axis.title = '⚡ Wave energy delivered each month', 'MJ per metre of wave crest'
    b2.add_data(Reference(wm, min_col=12, max_col=12, min_row=3, max_row=end_m), titles_from_data=True)
    b2.set_categories(Reference(wm, min_col=1, min_row=4, max_row=end_m))
    b2.width, b2.height = 20, 9
    b2.legend = None
    wm.add_chart(b2, f'L{tr + 3}')
    b3 = BarChart()
    b3.type = 'col'
    b3.title, b3.y_axis.title = '🌀 Storm hours per month', 'hours'
    b3.add_data(Reference(wm, min_col=22, max_col=22, min_row=3, max_row=end_m), titles_from_data=True)
    b3.set_categories(Reference(wm, min_col=1, min_row=4, max_row=end_m))
    b3.width, b3.height = 20, 9
    b3.legend = None
    wm.add_chart(b3, f'A{tr + 22}')

    # ================================================================ Storms ==
    ws = wb.create_sheet('Storms')
    _title(ws, '🌀 Storm events',
           f'Event = hourly Hs ≥ {bc.STORM_HS_M:g} m for {bc.STORM_MIN_HOURS}+ h (dips < {bc.STORM_MERGE_GAP_H} h merged). '
           'Settings live at the top of buoy_tools/buoy_core.py.', 12)
    _header(ws, 3, ['#', 'Start (local)', 'End (local)', 'Duration (h)', 'Peak Hs (m)', 'Peak Hs (ft)', 'Peak time',
                    'Tp at peak (s)', 'Direction', 'Max wind* (m/s)', 'Min pressure (hPa)', 'Energy (MJ/m)', 'Label'])
    end_s = 3
    if len(ev):
        e = ev[['event', 'start', 'end', 'hours', 'peak_hs', 'peak_hs', 'peak_time', 'tp_at_peak', 'dir', 'max_wind',
                'min_pres', 'energy_mj_m']].copy()
        e['label'] = ev['start'].dt.strftime("%b %d '%y")
        e.columns = range(e.shape[1])
        end_s = _table(ws, 4, e, ['0', 'yyyy-mm-dd hh:mm', 'yyyy-mm-dd hh:mm', '0', '0.00', '0.0', 'yyyy-mm-dd hh:mm',
                                  '0.0', '@', '0.0', '0.0', '#,##0', '@']) - 1
        for r in range(4, end_s + 1):          # feet as a live formula
            ws.cell(row=r, column=6, value=f'=E{r}*3.28084').number_format = '0.0'
        bs = BarChart()
        bs.type = 'col'
        bs.title, bs.y_axis.title = '🌀 Peak wave height of each storm', 'Hs (m)'
        bs.add_data(Reference(ws, min_col=5, max_col=5, min_row=3, max_row=end_s), titles_from_data=True)
        bs.set_categories(Reference(ws, min_col=13, min_row=4, max_row=end_s))
        bs.width, bs.height = 24, 9
        bs.legend = None
        ws.add_chart(bs, f'A{end_s + 3}')
    _widths(ws, [5, 17, 17, 10, 10, 10, 17, 10, 10, 11, 12, 11, 11])
    ws.freeze_panes = 'A4'

    # ================================================================== Rose ==
    wr = wb.create_sheet('Wave Rose')
    _title(wr, '🧭 Where do the waves come from?',
           '% of all wave readings, by direction the waves are coming FROM (rows) and wave height (columns).', 9)
    r = ro.reset_index().rename(columns={'dm': 'Direction', 'index': 'Direction'})
    r.columns = ['Direction'] + list(ro.columns)
    for c in ro.columns:
        r[c] = r[c] / 100
    _header(wr, 3, list(r.columns) + ['Total'])
    end_r = _table(wr, 4, r, ['@'] + ['0.0%'] * len(ro.columns)) - 1
    tc = len(ro.columns) + 2
    for rr in range(4, end_r + 1):
        c = wr.cell(row=rr, column=tc, value=f'=SUM(B{rr}:{get_column_letter(tc - 1)}{rr})')
        c.number_format, c.font, c.border = '0.0%', _font(bold=True), BOX
    wr.cell(row=end_r + 1, column=1, value='Total').font = _font(bold=True)
    for cc in range(2, tc + 1):
        L = get_column_letter(cc)
        c = wr.cell(row=end_r + 1, column=cc, value=f'=SUM({L}4:{L}{end_r})')
        c.number_format, c.font, c.border = '0.0%', _font(bold=True), BOX
    _widths(wr, [11] + [10] * (tc - 1))
    br = BarChart()
    br.type, br.grouping, br.overlap = 'col', 'stacked', 100
    br.title, br.y_axis.title = '🧭 Wave direction × height', 'share of readings'
    br.add_data(Reference(wr, min_col=2, max_col=tc - 1, min_row=3, max_row=end_r), titles_from_data=True)
    br.set_categories(Reference(wr, min_col=1, min_row=4, max_row=end_r))
    br.y_axis.number_format = '0%'
    br.width, br.height = 22, 10
    br.legend.position = 'b'
    wr.add_chart(br, f'A{end_r + 4}')

    # ================================================================ Health ==
    wk = wb.create_sheet('Buoy Health')
    _title(wk, '🔧 Buoy health', 'Monthly housekeeping. Rising hull humidity can mean moisture is getting inside.', 5)
    hk = S['monthly'][~S['monthly']['partial']][['label', 'batt_min', 'humid_mean', 'coverage_pct']].copy()
    hk['coverage_pct'] /= 100
    _header(wk, 3, ['Month', 'Battery min (V)', 'Hull humidity mean (%)', 'Data coverage'])
    end_k = _table(wk, 4, hk, ['@', '0.00', '0.0', '0%']) - 1
    _widths(wk, [17, 17, 14, 12])
    lk = _line_chart('💧 Hull humidity (monthly mean)', '% relative humidity', 18, 8)
    lk.add_data(Reference(wk, min_col=3, max_col=3, min_row=3, max_row=end_k), titles_from_data=True)
    lk.set_categories(Reference(wk, min_col=1, min_row=4, max_row=end_k))
    lk.legend = None
    wk.add_chart(lk, 'F3')
    rr = end_k + 3
    rr = _section(wk, rr, 'Data gaps longer than 3 hours', 4)
    _header(wk, rr, ['From', 'To', 'Hours', 'Days'])
    for i, (a, b_, hh) in enumerate(S['gaps'], start=rr + 1):
        for j, v in enumerate([a.tz_localize(None).to_pydatetime(), b_.tz_localize(None).to_pydatetime(), hh, f'=C{i}/24'], 1):
            c = wk.cell(row=i, column=j, value=v)
            c.font, c.border = _font(), BOX
            c.number_format = ['yyyy-mm-dd hh:mm', 'yyyy-mm-dd hh:mm', '0.0', '0.0'][j - 1]

    # ============================================================== Overview ==
    wo = wb.create_sheet('Overview', 0)
    _title(wo, f'🌊 Camp Ellis Wave Buoy ({st["spotter"]}) — Data Summary',
           f'{st["site"]} · {st["first"]:%b %d, %Y} → {st["last"]:%b %d, %Y} · '
           f'{st["n_files"]} monthly files · generated {dt.datetime.now():%b %d, %Y %I:%M %p}', 6)
    _widths(wo, [40, 16, 10, 14, 70, 4])
    row = 4
    row = _section(wo, row, '📊 Headline numbers  (live formulas over the Hourly / Daily / Monthly sheets)')
    _header(wo, row, ['Metric', 'Value', 'Unit', 'In feet / other', 'What it means'])
    row += 1
    B = HR('B')
    lines = [
        ('Days of data', f'=(MAX({HR("A")})-MIN({HR("A")}))', 'days', None,
         'Time between first and last reading.', '0.0'),
        ('Hours with wave data', f'=COUNT({B})', 'hours', f'=B{{r}}/(B{{r0}}*24)',
         'Right-hand column = share of the deployment covered.', '#,##0', '0.0%'),
        ('Average wave height (Hs)', f'=AVERAGE({B})', 'm', f'=B{{r}}*3.28084',
         'Significant wave height ≈ the height an observer would call "the waves".', '0.00', '0.0 "ft"'),
        ('Typical (median) wave height', f'=MEDIAN({B})', 'm', f'=B{{r}}*3.28084',
         'Half the time waves were smaller than this.', '0.00', '0.0 "ft"'),
        ('90th percentile wave height', f'=PERCENTILE({B},0.9)', 'm', f'=B{{r}}*3.28084',
         'Only 10% of hours were rougher than this.', '0.00', '0.0 "ft"'),
        ('Biggest hourly wave height', f'=MAX({B})', 'm', f'=B{{r}}*3.28084',
         'Highest hourly-average Hs (15-min peak is on the Storms sheet).', '0.00', '0.0 "ft"'),
        ('…which happened on', f'=INDEX({HR("A")},MATCH(MAX({B}),{B},0))', '', None,
         'Local time.', 'yyyy-mm-dd hh:mm'),
        ('Share of hours with Hs ≥ 1 m (3.3 ft)', f'=COUNTIF({B},">=1")/COUNT({B})', '', None,
         '"Rough" by this bay\'s standards.', '0.0%'),
        ('Share of hours calm (Hs < 0.3 m / 1 ft)', f'=COUNTIF({B},"<0.3")/COUNT({B})', '', None,
         'Glassy, lake-like conditions.', '0.0%'),
        ('Average peak period (Tp)', f'=AVERAGE({HR("C")})', 's', None,
         'Seconds between the dominant wave crests. Longer = more powerful swell.', '0.0'),
        ('Average wave power', f'=AVERAGE({HR("F")})', 'kW/m', None,
         'Energy flowing past each metre of wave crest (deep-water estimate).', '0.00'),
        ('Total wave energy', f'=SUM({HR("F")})/1000', 'MWh/m', f'=B{{r}}*1000/29',
         'Right-hand column ≈ days of electricity for a typical U.S. home (~29 kWh/day).', '0.0', '#,##0 "home-days"'),
        ('Lowest barometric pressure (hourly)', f'=MIN({HR("H")})', 'hPa', f'=B{{r}}*0.02953',
         'Storm-strength lows are below ~1000 hPa. Right column in inches of mercury.', '0.0', '0.00 "inHg"'),
        ('…which happened on', f'=INDEX({HR("A")},MATCH(MIN({HR("H")}),{HR("H")},0))', '', None, 'Local time.',
         'yyyy-mm-dd hh:mm'),
        ('Roughest day (highest daily mean Hs)', f'=INDEX(Daily!$A$4:$A${end_d},MATCH(MAX(Daily!$B$4:$B${end_d}),Daily!$B$4:$B${end_d},0))',
         '', f'=MAX(Daily!$B$4:$B${end_d})', 'Right-hand column = that day\'s mean Hs (m).', 'yyyy-mm-dd', '0.00 "m"'),
        ('Roughest month (highest mean Hs)', f'=INDEX(Monthly!$A$4:$A${end_m},MATCH(MAX(Monthly!$D$4:$D${end_m}),Monthly!$D$4:$D${end_m},0))',
         '', f'=MAX(Monthly!$D$4:$D${end_m})', 'Right-hand column = that month\'s mean Hs (m).', '@', '0.00 "m"'),
        ('Calmest month (lowest mean Hs)', f'=INDEX(Monthly!$A$4:$A${end_m},MATCH(MIN(Monthly!$D$4:$D${end_m}),Monthly!$D$4:$D${end_m},0))',
         '', f'=MIN(Monthly!$D$4:$D${end_m})', '', '@', '0.00 "m"'),
        ('Storm events', f'=COUNT(Storms!$A$4:$A${max(end_s, 4)})', 'events', f'=SUM(Storms!$D$4:$D${max(end_s, 4)})',
         f'Hs ≥ {bc.STORM_HS_M:g} m for {bc.STORM_MIN_HOURS}+ h. Right column = total storm hours.', '0', '0 "h"'),
        ('Most common wave direction', st['dir_mode'], '', st['dir_mode_pct'] / 100,
         'Direction waves come FROM. Right column = share of readings (see Wave Rose).', '@', '0%'),
    ]
    r0 = row
    for item in lines:
        label, val, unit, other, meaning, fmt = item[:6]
        fmt2 = item[6] if len(item) > 6 else fmt
        wo.cell(row=row, column=1, value=label).font = _font(bold=True)
        c = wo.cell(row=row, column=2, value=val)
        c.number_format, c.font = fmt, _font(color='000000')
        wo.cell(row=row, column=3, value=unit).font = _font(color='666666')
        if other is not None:
            o = wo.cell(row=row, column=4, value=other.format(r=row, r0=r0) if isinstance(other, str) else other)
            o.number_format, o.font = fmt2, _font(color='1D7A8C')
        wo.cell(row=row, column=5, value=meaning).font = _font(color='444444')
        wo.cell(row=row, column=5).alignment = Alignment(wrap_text=True, vertical='top')
        for j in range(1, 6):
            wo.cell(row=row, column=j).border = BOX
            if (row - r0) % 2:
                wo.cell(row=row, column=j).fill = ALT_FILL
        row += 1

    row += 1
    row = _section(wo, row, '🎛️ Try it yourself — change the yellow cell')
    wo.cell(row=row, column=1, value='Wave height threshold (m)').font = _font(bold=True)
    c = wo.cell(row=row, column=2, value=1.5)
    c.fill, c.font, c.number_format, c.border = IN_FILL, _font(bold=True, color='0000FF'), '0.00', BOX
    wo.cell(row=row, column=4, value=f'=B{row}*3.28084').number_format = '0.0 "ft"'
    wo.cell(row=row, column=5, value='Type any height in metres — the rows below update.').font = _font(italic=True, color='444444')
    thr = row
    row += 1
    for label, f_, fmt in [('Hours above that height', f'=COUNTIF({B},">="&B{thr})', '#,##0'),
                           ('Share of the time', f'=COUNTIF({B},">="&B{thr})/COUNT({B})', '0.00%'),
                           ('Days that reached it at least once', f'=COUNTIF(Daily!$C$4:$C${end_d},">="&B{thr})', '#,##0')]:
        wo.cell(row=row, column=1, value=label).font = _font()
        c = wo.cell(row=row, column=2, value=f_)
        c.number_format, c.border = fmt, BOX
        row += 1

    row += 1
    row = _section(wo, row, '💡 What the data says')
    import sensors as sn
    for emo, head, detail in bc.insights(S) + (sn.insights(SS, S) if SS else []):
        wo.cell(row=row, column=1, value=f'{emo}  {head}').font = _font(bold=True, color=NAVY)
        wo.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical='top')
        wo.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        c = wo.cell(row=row, column=2, value=detail)
        c.font = _font()
        c.alignment = Alignment(wrap_text=True, vertical='top')
        wo.row_dimensions[row].height = max(30, 14 * (len(detail) // 95 + 1))
        row += 1

    row += 1
    row = _section(wo, row, '⚠️ Data-quality notes (what was cleaned and why)')
    for n in bc.qc_notes(S):
        wo.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        c = wo.cell(row=row, column=1, value='•  ' + n)
        c.font = _font(color='444444')
        c.alignment = Alignment(wrap_text=True, vertical='top')
        wo.row_dimensions[row].height = max(16, 14 * (len(n) // 150 + 1))
        row += 1

    row += 1
    row = _section(wo, row, '📖 Glossary')
    gl = [('Hs — significant wave height', 'Average height (trough to crest) of the highest one-third of waves. Individual waves can be ~2× Hs.'),
          ('Tp — peak period', 'Time between crests of the waves carrying the most energy. >10 s = long-travelled swell; <6 s = local chop.'),
          ('Te — energy period', 'Period used to compute wave power; calculated from each wave spectrum.'),
          ('Direction', 'Compass direction the waves are coming FROM (0° = N, 90° = E).'),
          ('Swell vs. sea', 'Swell = long-period waves from distant storms (here: periods > 7.9 s); sea = short waves from local wind.'),
          ('Wave power (kW/m)', 'Energy per second crossing each metre of wave crest.'),
          ('* Wind', 'Estimated by the Spotter from the waves — not an anemometer measurement.')]
    for k, v in gl:
        wo.cell(row=row, column=1, value=k).font = _font(bold=True)
        wo.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        c = wo.cell(row=row, column=2, value=v)
        c.font = _font()
        c.alignment = Alignment(wrap_text=True)
        row += 1
    wo.cell(row=row + 1, column=1,
            value='Built automatically by buoy_tools/update_buoy_report.py — rerun it after adding new monthly CSVs.'
            ).font = _font(italic=True, size=8, color='888888')
    wo.freeze_panes = 'A4'

    _sensor_sheets(wb, SS)
    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
        w.sheet_properties.tabColor = {'Overview': NAVY, 'Monthly': TEAL, 'Storms': CORAL, 'Water Quality': '1BAF7A'}.get(w.title, '9FC5D0')
    wb.save(path)
    return path


def _sensor_sheets(wb, SS):
    """'Water Quality' (summary + monthly + charts) and 'Sensors Hourly' (raw hourly table)."""
    ws = wb.create_sheet('Water Quality', 1)
    _title(ws, '🧪 Below the surface — Smart Mooring sensors',
           'Subsurface water temperature, dissolved oxygen, currents and any other sensors on the mooring line.', 9)
    _widths(ws, [48, 12, 10, 11, 11, 11, 11, 11, 30])
    if not SS or not SS.get('has'):
        ws['A4'] = '🔌 No Smart Mooring sensor data has been received yet — this sheet fills in automatically once it arrives.'
        ws['A4'].font = _font(italic=True, color='444444')
        return
    _header(ws, 3, ['Channel', 'Position', 'Units', 'Latest', 'Average', 'Min', 'Max', 'Readings', 'Sofar data-type name'])
    rows = pd.DataFrame([[s['col'], s['pos_txt'], s['units'], s['latest'], s['mean'], s['min'], s['max'], s['n'], s['dtype']]
                         for s in SS['series']])
    end = _table(ws, 4, rows, ['@', '@', '@', '0.00', '0.00', '0.00', '0.00', '#,##0', '@']) - 1
    r = end + 2
    r = _section(ws, r, '📅 Monthly averages (hourly data)', 9)
    mo = SS['monthly']
    if len(mo):
        piv = mo.pivot_table(index='month', columns='series', values='mean', aggfunc='first')
        order = sorted(piv.index, key=lambda m: pd.Timestamp('1 ' + m))
        piv = piv.reindex(order)[[s['col'] for s in SS['series'] if s['col'] in piv.columns]]
        _header(ws, r, ['Month'] + list(piv.columns))
        start = r + 1
        endm = _table(ws, start, piv.reset_index(), ['@'] + ['0.00'] * len(piv.columns)) - 1
        for j in range(2, len(piv.columns) + 2):
            ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, 14)
        r = endm + 2
    # hourly sheet + charts per group
    wh = wb.create_sheet('Sensors Hourly')
    H = SS['hourly'].copy()
    H.index = bc.local(H.index).tz_localize(None)
    H = H.dropna(how='all')
    _title(wh, '⏱️ Smart Mooring sensors — hourly averages', 'Local time. One column per sensor channel.', len(H.columns) + 1)
    _header(wh, 3, ['Time (local)'] + list(H.columns))
    endh = _table(wh, 4, H.reset_index(), ['yyyy-mm-dd hh:mm'] + ['0.000'] * len(H.columns), band=False) - 1
    _widths(wh, [18] + [16] * len(H.columns))
    wh.freeze_panes = 'B4'
    groups = []
    for s in SS['series']:
        if s['group'] not in groups and s['group'] not in ('cur_dir', 'diag'):
            groups.append(s['group'])
    anchor = r + 1
    for g in groups[:6]:
        idx = [i for i, s in enumerate(SS['series']) if s['group'] == g]
        ch = _line_chart(f"{SS['series'][idx[0]]['emoji']} {SS['series'][idx[0]]['name']}", SS['series'][idx[0]]['units'], 22, 8)
        for i in idx:
            ch.add_data(Reference(wh, min_col=i + 2, max_col=i + 2, min_row=3, max_row=endh), titles_from_data=True)
        ch.set_categories(Reference(wh, min_col=1, min_row=4, max_row=endh))
        ch.x_axis.number_format = 'mmm yy'
        ws.add_chart(ch, f'A{anchor}')
        anchor += 18
